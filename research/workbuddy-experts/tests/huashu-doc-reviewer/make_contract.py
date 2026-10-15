"""生成一份埋了 6 个已知问题的测试合同，用于检验「花叔文档审稿专家」的召回率。"""
from docx import Document
from docx.shared import Pt

PARAS = [
    ("title", "软件开发服务合同"),
    ("p", "甲方：上海某某贸易有限公司    乙方：杭州某某科技有限公司"),
    ("h", "第一条 合同金额"),
    ("p", "本项目合同总价为人民币50万元，甲方根据本合同约定向乙方支付。"),          # 问题1 含税口径缺失；问题6 "根据"→"依据"
    ("h", "第二条 付款方式"),
    ("p", "甲方在收到乙方开具的发票后尽快付款。"),                                 # 问题2 账期模糊
    ("h", "第三条 服务标准"),
    ("p", "乙方应保证系统稳定运行，出现故障时及时处理。"),                         # 问题3 SLA 未量化
    ("h", "第四条 保密"),
    ("p", "双方应对合作中知悉的对方信息保密。"),                                   # 问题4 保密范围/期限缺失
    ("h", "第五条 争议解决"),
    ("p", "如有争议，双方协商解决，协商不成的，任何一方可向法院起诉。"),           # 问题5 管辖法院不明
    ("h", "第六条 其他"),
    ("p", "本合同一式两份，双方各执一份，自双方签字盖章之日起生效。"),
]

doc = Document()
style = doc.styles["Normal"]; style.font.name = "SimSun"; style.font.size = Pt(11)
for kind, text in PARAS:
    if kind == "title": doc.add_heading(text, level=0)
    elif kind == "h": doc.add_heading(text, level=2)
    else: doc.add_paragraph(text)
doc.save("contract.docx")
print("contract.docx written")
